"""Extract text from PDF, PPTX, DOCX, TXT, and images (OCR when needed)."""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from pathlib import Path

logger = logging.getLogger(__name__)


@dataclass
class ExtractedPage:
    page_number: int
    heading: str
    text: str
    method: str
    ocr_confidence: float | None = None


@dataclass
class ExtractionResult:
    pages: list[ExtractedPage] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _ocr_image(path: Path) -> tuple[str, float | None]:
    try:
        import pytesseract
        from PIL import Image

        img = Image.open(path)
        data = pytesseract.image_to_data(img, output_type=pytesseract.Output.DICT)
        text = pytesseract.image_to_string(img)
        confs = [float(c) for c in data.get("conf", []) if str(c) not in ("-1", "")]
        conf = sum(confs) / len(confs) if confs else None
        return text, conf
    except Exception as exc:  # OCR is optional if tesseract is missing
        logger.warning("OCR unavailable: %s", exc)
        return "", None


def extract_pdf(path: Path) -> ExtractionResult:
    import fitz

    result = ExtractionResult()
    doc = fitz.open(path)
    try:
        for i, page in enumerate(doc, start=1):
            text = page.get_text("text") or ""
            heading = ""
            blocks = page.get_text("dict").get("blocks", [])
            for block in blocks:
                for line in block.get("lines", []):
                    spans = line.get("spans", [])
                    if not spans:
                        continue
                    size = spans[0].get("size", 0)
                    candidate = "".join(s.get("text", "") for s in spans).strip()
                    if size >= 14 and candidate:
                        heading = candidate[:200]
                        break
                if heading:
                    break
            method = "pymupdf"
            conf = None
            if not text.strip():
                pix = page.get_pixmap(dpi=200)
                tmp = path.parent / f".ocr_page_{i}.png"
                pix.save(tmp)
                ocr_text, conf = _ocr_image(tmp)
                tmp.unlink(missing_ok=True)
                text = ocr_text
                method = "tesseract"
            result.pages.append(
                ExtractedPage(page_number=i, heading=heading, text=text, method=method, ocr_confidence=conf)
            )
    finally:
        doc.close()
    return result


def extract_pptx(path: Path) -> ExtractionResult:
    from pptx import Presentation

    result = ExtractionResult()
    prs = Presentation(str(path))
    for i, slide in enumerate(prs.slides, start=1):
        texts: list[str] = []
        heading = ""
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            t = shape.text_frame.text.strip()
            if not t:
                continue
            if not heading:
                heading = t.splitlines()[0][:200]
            texts.append(t)
        result.pages.append(
            ExtractedPage(page_number=i, heading=heading, text="\n".join(texts), method="python-pptx")
        )
    return result


def extract_docx(path: Path) -> ExtractionResult:
    from docx import Document as Docx

    result = ExtractionResult()
    doc = Docx(str(path))
    pages: list[list[str]] = [[]]
    heading = ""
    page_no = 1
    current_heading = ""
    for para in doc.paragraphs:
        style = (para.style.name or "").lower() if para.style else ""
        text = para.text.strip()
        if not text:
            continue
        if "heading" in style:
            if pages[-1]:
                result.pages.append(
                    ExtractedPage(
                        page_number=page_no, heading=current_heading, text="\n".join(pages[-1]), method="python-docx"
                    )
                )
                page_no += 1
                pages.append([])
            current_heading = text[:200]
            if not heading:
                heading = current_heading
        pages[-1].append(text)
    if pages[-1]:
        result.pages.append(
            ExtractedPage(
                page_number=page_no, heading=current_heading or heading, text="\n".join(pages[-1]), method="python-docx"
            )
        )
    if not result.pages:
        result.pages.append(ExtractedPage(page_number=1, heading="", text="", method="python-docx"))
    return result


def extract_txt(path: Path) -> ExtractionResult:
    text = path.read_text(encoding="utf-8", errors="replace")
    heading = text.splitlines()[0][:200] if text.strip() else ""
    return ExtractionResult(pages=[ExtractedPage(page_number=1, heading=heading, text=text, method="txt")])


def extract_image(path: Path) -> ExtractionResult:
    text, conf = _ocr_image(path)
    return ExtractionResult(
        pages=[ExtractedPage(page_number=1, heading="", text=text, method="tesseract", ocr_confidence=conf)]
    )


def extract_document(path: Path) -> ExtractionResult:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return extract_pdf(path)
    if ext == ".pptx":
        return extract_pptx(path)
    if ext == ".docx":
        return extract_docx(path)
    if ext == ".txt":
        return extract_txt(path)
    if ext in {".png", ".jpg", ".jpeg"}:
        return extract_image(path)
    raise ValueError(f"Unsupported extraction type: {ext}")
